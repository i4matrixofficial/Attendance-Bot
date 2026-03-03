import os
from datetime import datetime
import discord
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

DATA_DIR = "data"

# ── Color palette ──────────────────────────────────────────────────────────────
COLOR_HEADER_BG   = "1A2B4A"   # Dark navy  – title row
COLOR_HEADER_FG   = "FFFFFF"   # White      – title text
COLOR_META_BG     = "2C3E6B"   # Mid navy   – info rows
COLOR_META_FG     = "FFFFFF"
COLOR_COL_HDR_BG  = "3A5BA0"   # Blue       – column header row
COLOR_COL_HDR_FG  = "FFFFFF"
COLOR_PRESENT_BG  = "D6F5DD"   # Light green – present rows
COLOR_PRESENT_FG  = "1A6B2E"   # Dark green  – present text
COLOR_ABSENT_BG   = "FCE8E8"   # Light red   – absent rows
COLOR_ABSENT_FG   = "9B1C1C"   # Dark red    – absent text
COLOR_ALT_PRESENT = "C2EEC9"   # Slightly darker green for alternating
COLOR_ALT_ABSENT  = "F8D7D7"   # Slightly darker red   for alternating
COLOR_BORDER      = "AAAAAA"

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(hex_color, bold=False, size=11):
    return Font(color=hex_color, bold=bold, size=size, name="Calibri")

def _border():
    thin = Side(style="thin", color=COLOR_BORDER)
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center")

# ── Main function ──────────────────────────────────────────────────────────────
def record_attendance(channel: discord.VoiceChannel) -> str:
    """
    Records attendance for all non-bot guild members.
    Saves a colored Excel (.xlsx) report and returns the file path.
    """
    if not os.path.exists(DATA_DIR):
        os.makedirs(DATA_DIR)

    now          = datetime.now()
    date_str     = now.strftime("%Y-%m-%d")
    time_str     = now.strftime("%H:%M:%S")
    timestamp_str = now.strftime("%Y%m%d_%H%M%S")

    safe_ch = "".join(c for c in channel.name if c.isalnum() or c == " ").strip()
    filename  = f"Attendance_Report_{safe_ch}_{timestamp_str}.xlsx"
    filepath  = os.path.join(DATA_DIR, filename)

    # Members present in voice channel (exclude bots)
    present_ids = {m.id for m in channel.members if not m.bot}

    # All guild members excluding bots
    all_humans = [m for m in channel.guild.members if not m.bot]
    all_humans.sort(key=lambda m: (m.id not in present_ids, m.display_name.lower()))

    present_count = len(present_ids)
    absent_count  = len(all_humans) - present_count

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # ── Column widths ──────────────────────────────────────────────────────────
    col_widths = [5, 16, 22, 26, 22]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Helper: apply style to a row ──────────────────────────────────────────
    def style_row(row_num, bg, fg, bold=False, align_fn=_left):
        for cell in ws[row_num]:
            cell.fill      = _fill(bg)
            cell.font      = _font(fg, bold=bold)
            cell.alignment = align_fn()
            cell.border    = _border()

    # ── Row 1: Company title ───────────────────────────────────────────────────
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A1:E1")
    ws["A1"] = "🏢  i4matrix Attendance Report"
    ws["A1"].fill      = _fill(COLOR_HEADER_BG)
    ws["A1"].font      = Font(color=COLOR_HEADER_FG, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = _center()
    ws["A1"].border    = _border()

    # ── Rows 2-6: Metadata ────────────────────────────────────────────────────
    meta_rows = [
        ("Server",          channel.guild.name),
        ("Voice Channel",   channel.name),
        ("Date",            date_str),
        ("Time",            time_str),
        ("Present / Total", f"{present_count} Present  |  {absent_count} Absent  |  {len(all_humans)} Total"),
    ]
    for i, (label, value) in enumerate(meta_rows, start=2):
        ws.row_dimensions[i].height = 18
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value
        ws.merge_cells(f"B{i}:E{i}")
        style_row(i, COLOR_META_BG, COLOR_META_FG, bold=(i == 2), align_fn=_left)
        ws[f"A{i}"].font = Font(color=COLOR_META_FG, bold=True, size=10, name="Calibri")

    # ── Row 7: Column headers ─────────────────────────────────────────────────
    HDR_ROW = 7
    ws.row_dimensions[HDR_ROW].height = 22
    headers = ["#", "Status", "Display Name", "Username", "User ID"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=HDR_ROW, column=col, value=h)
        cell.fill      = _fill(COLOR_COL_HDR_BG)
        cell.font      = _font(COLOR_COL_HDR_FG, bold=True)
        cell.alignment = _center()
        cell.border    = _border()

    # ── Data rows ──────────────────────────────────────────────────────────────
    for idx, member in enumerate(all_humans, start=1):
        row_num = HDR_ROW + idx
        ws.row_dimensions[row_num].height = 18
        is_present  = member.id in present_ids
        # Alternate shading so rows are easier to scan
        if is_present:
            bg = COLOR_PRESENT_BG if idx % 2 == 1 else COLOR_ALT_PRESENT
            fg = COLOR_PRESENT_FG
            status = "✅ Present"
        else:
            bg = COLOR_ABSENT_BG if idx % 2 == 1 else COLOR_ALT_ABSENT
            fg = COLOR_ABSENT_FG
            status = "❌ Absent"

        row_data = [idx, status, member.display_name, f"@{member.name}", str(member.id)]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.fill      = _fill(bg)
            cell.font      = _font(fg, bold=is_present)
            cell.alignment = _center() if col in (1, 2) else _left()
            cell.border    = _border()

    # ── Freeze header rows so they stay visible when scrolling ────────────────
    ws.freeze_panes = f"A{HDR_ROW + 1}"

    wb.save(filepath)
    return filepath
