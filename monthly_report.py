import calendar
import os
from datetime import date, datetime, timedelta
from typing import Optional, Sequence

import discord
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from attendance_db import (
    get_day_records,
    get_logged_dates_range,
    get_month_records,
    get_period_records,
)

COLOR_HEADER_BG = "1A2B4A"
COLOR_HEADER_FG = "FFFFFF"
COLOR_META_BG = "2C3E6B"
COLOR_META_FG = "FFFFFF"
COLOR_COL_HDR_BG = "3A5BA0"
COLOR_COL_HDR_FG = "FFFFFF"
COLOR_GREEN = "D6F5DD"
COLOR_YELLOW = "FDF1CC"
COLOR_RED = "FCE8E8"
COLOR_BORDER = "AAAAAA"


def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)


def _font(hex_color: str, bold: bool = False, size: int = 11):
    return Font(color=hex_color, bold=bold, size=size, name="Calibri")


def _border():
    thin = Side(style="thin", color=COLOR_BORDER)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left():
    return Alignment(horizontal="left", vertical="center")


def _tier_color(percentage: float) -> str:
    if percentage >= 90:
        return COLOR_GREEN
    if percentage >= 70:
        return COLOR_YELLOW
    return COLOR_RED


def _safe_channel_name(channel_name: str) -> str:
    return "".join(
        c for c in channel_name if c.isalnum() or c in (" ", "_", "-")
    ).strip().replace(" ", "_")


def _make_workbook(
    title: str,
    meta: Sequence[tuple[str, str]],
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    filepath: str,
) -> str:
    if not os.path.exists("data"):
        os.makedirs("data")

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"

    widths = [5, 34, 16, 16, 16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.merge_cells("A1:E1")
    ws["A1"] = title
    ws["A1"].fill = _fill(COLOR_HEADER_BG)
    ws["A1"].font = _font(COLOR_HEADER_FG, bold=True, size=16)
    ws["A1"].alignment = _center()
    ws["A1"].border = _border()

    for idx, (label, value) in enumerate(meta, start=2):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = value
        ws.merge_cells(f"B{idx}:E{idx}")
        ws[f"A{idx}"].fill = _fill(COLOR_META_BG)
        ws[f"A{idx}"].font = _font(COLOR_META_FG, bold=True, size=10)
        ws[f"A{idx}"].alignment = _left()
        ws[f"A{idx}"].border = _border()
        ws[f"B{idx}"].fill = _fill(COLOR_META_BG)
        ws[f"B{idx}"].font = _font(COLOR_META_FG, size=10)
        ws[f"B{idx}"].alignment = _left()
        ws[f"B{idx}"].border = _border()

    header_row = 6
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = _fill(COLOR_COL_HDR_BG)
        cell.font = _font(COLOR_COL_HDR_FG, bold=True)
        cell.alignment = _center()
        cell.border = _border()

    for row_idx, row_values in enumerate(rows, start=1):
        row = header_row + row_idx
        for col, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.fill = _fill(_tier_color(float(value[:-1])) if col == 5 and isinstance(value, str) and value.endswith('%') else COLOR_GREEN if col != 5 else COLOR_GREEN)
            cell.font = _font("000000", bold=(col == 2))
            cell.alignment = _center() if col != 2 else _left()
            cell.border = _border()

    ws.freeze_panes = f"A{header_row + 1}"
    wb.save(filepath)
    return filepath


def _compute_percentage(present: int, total: int) -> str:
    return f"{(present / total * 100) if total else 0:.0f}%"


async def generate_daily_report(
    guild: discord.Guild,
    channel: discord.VoiceChannel,
    target_date: date,
) -> Optional[str]:
    records = await get_day_records(channel.id, target_date)
    if not records:
        return None

    safe_channel = _safe_channel_name(channel.name)
    filename = f"Daily_Attendance_{safe_channel}_{target_date.isoformat()}.xlsx"
    filepath = os.path.join("data", filename)

    rows = []
    for idx, stats in enumerate(records.values(), start=1):
        rows.append([
            idx,
            stats["name"],
            stats["status"].capitalize(),
            "",
            "",
        ])

    return _make_workbook(
        title=f"🏢 i4matrix Daily Attendance — {target_date.isoformat()}",
        meta=[
            ("Server", guild.name),
            ("Voice Channel", channel.name),
            ("Date", target_date.isoformat()),
            ("Logged Members", str(len(records))),
        ],
        headers=["#", "Member", "Status", "", ""],
        rows=rows,
        filepath=filepath,
    )


async def generate_weekly_report(
    guild: discord.Guild,
    channel: discord.VoiceChannel,
    year: int,
    week_number: int,
) -> Optional[str]:
    start_date = date.fromisocalendar(year, week_number, 1)
    end_date = start_date + timedelta(days=6)
    records = await get_period_records(channel.id, start_date, end_date)
    logged_dates = await get_logged_dates_range(channel.id, start_date, end_date)

    if not records or not logged_dates:
        return None

    safe_channel = _safe_channel_name(channel.name)
    filename = f"Weekly_Attendance_{safe_channel}_{year}-W{week_number:02d}.xlsx"
    filepath = os.path.join("data", filename)

    rows = []
    for idx, stats in enumerate(records.values(), start=1):
        rows.append([
            idx,
            stats["name"],
            str(stats["present"]),
            str(stats["total"]),
            _compute_percentage(stats["present"], stats["total"]),
        ])

    return _make_workbook(
        title=f"🏢 i4matrix Weekly Attendance — {year} W{week_number:02d}",
        meta=[
            ("Server", guild.name),
            ("Voice Channel", channel.name),
            ("Week", f"{year}-W{week_number:02d}"),
            ("Days Tracked", str(len(logged_dates))),
        ],
        headers=["#", "Member", "Days Present", "Days Tracked", "Attendance %"],
        rows=rows,
        filepath=filepath,
    )


async def generate_monthly_report(
    guild: discord.Guild,
    channel: discord.VoiceChannel,
    year: int,
    month: int,
) -> Optional[str]:
    records = await get_month_records(channel.id, year, month)
    last_day = calendar.monthrange(year, month)[1]
    logged_dates = await get_logged_dates_range(channel.id, date(year, month, 1), date(year, month, last_day))

    if not records or not logged_dates:
        return None

    safe_channel = _safe_channel_name(channel.name)
    filename = f"Monthly_Attendance_{safe_channel}_{year}-{month:02d}.xlsx"
    filepath = os.path.join("data", filename)

    rows = []
    for idx, stats in enumerate(records.values(), start=1):
        rows.append([
            idx,
            stats["name"],
            str(stats["present"]),
            str(stats["total"]),
            _compute_percentage(stats["present"], stats["total"]),
        ])

    return _make_workbook(
        title=f"🏢 i4matrix Monthly Attendance Summary — {datetime(year, month, 1).strftime('%B %Y')}",
        meta=[
            ("Server", guild.name),
            ("Voice Channel", channel.name),
            ("Month", datetime(year, month, 1).strftime("%B %Y")),
            ("Days Tracked", str(len(logged_dates))),
        ],
        headers=["#", "Member", "Days Present", "Days Tracked", "Attendance %"],
        rows=rows,
        filepath=filepath,
    )
